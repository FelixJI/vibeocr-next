"""Retired seal history is projected without altering stored mixed documents."""

from __future__ import annotations

from copy import deepcopy

import pytest
from vibeocr.runtime.documents.export_service import ExportService
from vibeocr.runtime.recognition.models.ocr_result_serializer import (
    ocr_result_from_payload,
    project_retired_seals,
)


def mixed_history():
    return {
        "pipeline_type": "PP-StructureV3",
        "raw_text": "STALE_SEAL",
        "markdown_text": "STALE_SEAL",
        "html_text": "STALE_SEAL",
        "content_list": [
            {"type": "text", "text": "正文"},
            {"type": "seal", "text": "STALE_SEAL", "image": {"asset_id": "old-seal"}},
            {
                "type": "table",
                "table_body": "<table><tr><td colspan='2'>合并</td></tr><tr><td>A</td><td>B</td></tr></table>",
            },
            {"type": "formula", "text": "x+y"},
            {"type": "chart", "text": "历史图表"},
            {"type": "image", "text": "图片说明", "img_path": "kept-figure.png"},
            {"type": "list", "list_items": ["列表甲", "列表乙"]},
            {"type": "code", "code_body": "print(123)"},
            {"type": "footer", "text": "PAGE END"},
        ],
        "text_blocks": [{"label": "seal", "text": "STALE_SEAL", "content_index": 1}],
    }


def test_mixed_history_read_is_repeatable_and_preserves_nonseal_semantics():
    original = mixed_history()
    saved = deepcopy(original)
    first = project_retired_seals(original)
    assert first == project_retired_seals(original) == project_retired_seals(first)
    assert original == saved
    assert first["content_list"][1] == {"type": "retired"}
    assert first["content_list"][2]["type"] == "table"
    for key in ("raw_text", "markdown_text", "html_text"):
        assert "STALE_SEAL" not in first[key]
        for expected in (
            "正文",
            "合并",
            "x+y",
            "历史图表",
            "图片说明",
            "列表甲",
            "列表乙",
            "print(123)",
            "PAGE END",
        ):
            assert expected in first[key], (key, expected, first[key])
    assert "kept-figure.png" in first["markdown_text"]
    assert "colspan" in first["html_text"]
    assert "$$x+y$$" in first["markdown_text"]


@pytest.mark.parametrize("fmt", ["txt", "markdown", "html", "xlsx", "docx"])
def test_direct_history_export_filters_seal_without_mutating_source(tmp_path, fmt):
    original = mixed_history()
    result = ocr_result_from_payload(original)
    before = deepcopy(result)
    path = tmp_path / ("export." + fmt)
    assert ExportService.export(result, path, fmt)
    assert result == before
    if fmt in {"txt", "markdown", "html"}:
        text = path.read_text(encoding="utf-8")
    elif fmt == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path)
        text = "\n".join(
            str(cell.value) for sheet in workbook for row in sheet for cell in row
        )
        workbook.close()
    else:
        from docx import Document

        document = Document(path)
        text = "\n".join(p.text for p in document.paragraphs)
    assert "STALE_SEAL" not in text
    assert "正文" in text


def test_seal_only_history_is_empty_and_untyped_history_remains_unchanged():
    payload = {
        "raw_text": "STALE_SEAL",
        "markdown_text": "STALE_SEAL",
        "html_text": "STALE_SEAL",
        "content_list": [{"type": "seal", "text": "STALE_SEAL"}],
    }
    projected = project_retired_seals(payload)
    assert all(
        projected[key] == "" for key in ("raw_text", "markdown_text", "html_text")
    )
    flat = {"raw_text": "unclassified historical summary"}
    assert project_retired_seals(flat) is flat


def test_same_text_keeps_ordinary_line_confidence_and_content_only_seal_has_no_derived_text():
    payload = {
        "content_list": [
            {"type": "text", "text": "same"},
            {"type": "SEAL", "text": "same"},
        ],
        "text_blocks": [
            {"label": "text", "text": "same", "score": 0.7, "content_index": 0},
            {"label": "text", "text": "same", "score": 0.5, "content_index": 1},
        ],
        "text_with_scores": [["same", 0.7], ["same", 0.5]],
    }
    projected = project_retired_seals(payload)
    assert projected["raw_text"] == "same"
    assert projected["text_with_scores"] == [["same", 0.7]]
    assert projected["low_confidence_items"] == [["same", 0.7]]
    assert projected["avg_score"] == 0.7
    only = {
        "content_list": [{"type": "seal", "text": "private"}],
        "text_with_scores": [["private", 0.5]],
    }
    assert project_retired_seals(only)["text_with_scores"] == []


def test_cancelled_projection_never_returns_partially_filtered_history(monkeypatch):
    from vibeocr.runtime.documents.tables import reducer

    payload = mixed_history()
    saved = deepcopy(payload)
    monkeypatch.setattr(
        reducer, "build_result_projections", lambda _result, **_kwargs: None
    )
    with pytest.raises(RuntimeError, match="cancelled"):
        project_retired_seals(payload)
    assert payload == saved


def test_typed_seal_id_removes_associated_line_but_preserves_shared_ordinary_image():
    payload = {
        "content_list": [
            {
                "type": "seal",
                "block_id": "retired-id",
                "text": "seal",
                "img_path": "shared.png",
            },
            {"type": "image", "text": "ordinary", "img_path": "shared.png"},
        ],
        "text_blocks": [
            {"text": "seal", "label": "text", "content_id": "retired-id", "score": 0.9}
        ],
        "images": {"shared.png": {"present": True}},
    }
    result = project_retired_seals(payload)
    assert "seal" not in result["raw_text"]
    assert "ordinary" in result["raw_text"]
    assert result["images"] == payload["images"]
    assert result["text_with_scores"] == []


def test_legacy_history_without_pipeline_tag_preserves_every_nonseal_block():
    payload = mixed_history()
    payload.pop("pipeline_type")
    visible = project_retired_seals(payload)
    for key in ("raw_text", "markdown_text", "html_text"):
        assert "PAGE END" in visible[key]
        assert "STALE_SEAL" not in visible[key]
