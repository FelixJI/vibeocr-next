"""MinerU 4 transport contracts and real 4.0.2 result projections."""

import io
import json
import zipfile
from pathlib import Path

import httpx
import pytest
from vibeocr.runtime.recognition.mineru_api import (
    MineruApiClient,
    MineruApiError,
    MineruCancelled,
    MineruDocument,
)
from vibeocr.runtime.recognition.mineru_config import (
    MineruConfigError,
    resolve_mineru_config,
)
from vibeocr.runtime.recognition.mineru_result import project_document
from vibeocr.runtime_contracts import (
    ErrorCode,
    MineruConfig,
    MineruTier,
    PipelineSelection,
)

FIXTURES = Path(__file__).parent / "fixtures" / "mineru4"


def test_real_basic_three_page_projection():
    # Captured from unmodified MinerU4.0.2, basic/onnx, generated 3-page PDF.
    # These short lines are classified as headers upstream, intentionally
    # excluded from text by the existing product contract, retained in blocks.
    middle = json.loads((FIXTURES / "basic-middle.json").read_text())
    structured = json.loads((FIXTURES / "basic-structured.json").read_text())
    result = project_document(MineruDocument("", structured, middle, b""))
    assert [b["page_idx"] for b in result.content_list] == [0, 1, 2]
    assert result.content_list[0]["bbox"] == [115.0, 131.0, 364.0, 150.0]
    assert "page 3" in result.content_list[2]["text"]
    assert result.raw_text == ""


@pytest.mark.parametrize(
    "options",
    [
        {"backend": "pipeline"},
        {"backend": "vlm-engine"},
        {"enable_table": False},
        {"enable_formula": False},
        {"lang_list": ["ch", "korean"]},
        {"effort": "xhigh"},
    ],
)
def test_unequal_legacy_choices_require_reselection(options):
    with pytest.raises(MineruConfigError) as caught:
        resolve_mineru_config(PipelineSelection("MinerU", options=options))
    assert caught.value.code == ErrorCode.MINERU_CONFIG_MIGRATION_REQUIRED


def test_legacy_equivalence_and_explicit_tier():
    assert resolve_mineru_config(PipelineSelection("MinerU")).tier == MineruTier.BASIC
    migrated = resolve_mineru_config(
        PipelineSelection("MinerU", options={"effort": "high", "start_page_id": 2})
    )
    assert migrated.tier == MineruTier.STANDARD
    assert migrated.page_range == "3-r1"
    config = MineruConfig(tier=MineruTier.ADVANCED, language="korean")
    assert resolve_mineru_config(PipelineSelection("MinerU", mineru=config)) is config


class Server:
    def __init__(self, status="completed"):
        self.status = status
        self.calls = []
        self.job_payload = None
        self.middle = {
            "schema": "docvortex.middle",
            "schema_version": "2.0",
            "pages": [],
        }
        self.artifacts = {
            "markdown": b"",
            "middle_json": json.dumps(self.middle).encode(),
            "structured_content": b'{"pages": []}',
            "zip": b"zip",
        }

    def __call__(self, req):
        path = req.url.path
        self.calls.append((req.method, path))
        if req.method == "DELETE" or path.endswith("/cancel"):
            return httpx.Response(200, json={})
        if path == "/v1/uploads":
            return httpx.Response(
                200,
                json={
                    "id": "upload-1",
                    "upload_url": "https://untrusted.example/upload",
                },
            )
        if path == "/v1/uploads/upload-1/content":
            assert req.content == b"pdf"
            return httpx.Response(200, json={})
        if path == "/v1/uploads/upload-1/complete":
            return httpx.Response(200, json={"file": {"id": "input-1"}})
        if path == "/v1/parse/jobs":
            self.job_payload = json.loads(req.content)
            return httpx.Response(200, json={"job_id": "job-1", "status": "queued"})
        if path == "/v1/parse/jobs/job-1":
            return httpx.Response(
                200,
                json={
                    "job_id": "job-1",
                    "status": self.status,
                    "tier": self.job_payload["tier"],
                    "files": [
                        {
                            "name": "sample.pdf",
                            "status": "completed"
                            if self.status == "completed"
                            else "failed",
                            "error": {"code": "parse_failed"},
                            "output_files": {
                                k: {"file_id": k, "bytes": len(v)}
                                for k, v in self.artifacts.items()
                            },
                        }
                    ],
                },
            )
        if path.startswith("/v1/files/"):
            return httpx.Response(200, content=self.artifacts[path.split("/")[3]])
        raise AssertionError(path)


def test_upload_job_outputs_and_full_document(monkeypatch):
    monkeypatch.setattr(
        "vibeocr.runtime.recognition.mineru_api.time.sleep", lambda _: None
    )
    server = Server()
    result = MineruApiClient(
        "http://127.0.0.1", transport=httpx.MockTransport(server)
    ).parse([("sample.pdf", b"pdf")], MineruConfig(tier=MineruTier.BASIC))
    assert isinstance(result["sample.pdf"], MineruDocument)
    assert server.job_payload["files"][0]["page_range"] == "all"
    assert server.job_payload["tier"] == "basic"
    assert ("DELETE", "/v1/files/input-1") in server.calls
    assert ("DELETE", "/v1/files/zip") in server.calls
    assert all("file_parse" not in path for _, path in server.calls)


def test_partial_is_per_file_failure_not_success(monkeypatch):
    monkeypatch.setattr(
        "vibeocr.runtime.recognition.mineru_api.time.sleep", lambda _: None
    )
    server = Server("partial")
    result = MineruApiClient(
        "http://127.0.0.1", transport=httpx.MockTransport(server)
    ).parse([("sample.pdf", b"pdf")], MineruConfig(tier=MineruTier.BASIC))
    assert isinstance(result["sample.pdf"], MineruApiError)


def test_cancel_active_job_does_not_retry_or_change_tier(monkeypatch):
    monkeypatch.setattr(
        "vibeocr.runtime.recognition.mineru_api.time.sleep", lambda _: None
    )
    server = Server("running")

    def cancelled():
        return ("GET", "/v1/parse/jobs/job-1") in server.calls

    with pytest.raises(MineruCancelled):
        MineruApiClient(
            "http://127.0.0.1", transport=httpx.MockTransport(server)
        ).parse(
            [("sample.pdf", b"pdf")],
            MineruConfig(tier=MineruTier.ADVANCED),
            cancelled=cancelled,
        )
    assert ("DELETE", "/v1/parse/jobs/job-1") in server.calls
    assert server.calls.count(("POST", "/v1/parse/jobs")) == 1
    assert server.job_payload["tier"] == "advanced"


def test_unsafe_referenced_sidecar_is_rejected():
    middle = {
        "schema": "docvortex.middle",
        "schema_version": "2.0",
        "pages": [],
        "image_path": "../outside.png",
    }
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr("../outside.png", b"bad")
    with pytest.raises(MineruApiError, match="Unsafe"):
        project_document(MineruDocument("", {"pages": []}, middle, archive.getvalue()))


def test_real_native_table_and_embedded_image():
    middle = json.loads(
        (FIXTURES / "table-image-middle.json").read_text(encoding="utf8")
    )
    structured = json.loads(
        (FIXTURES / "table-image-structured.json").read_text(encoding="utf8")
    )
    result = project_document(MineruDocument("", structured, middle, b""))
    table = next(b for b in result.content_list if b["type"] == "table")
    assert "Apples" in table["text"] and "18" in table["text"]
    assert table["table"]["row_count"] == 3
    image = next(b for b in result.content_list if b["type"] == "image")
    assert result.images[image["img_path"]].startswith(b"\xff\xd8")
    assert "blue product sample" in result.raw_text


def test_unconfirmed_cancel_retains_inputs(monkeypatch):
    monkeypatch.setattr(
        "vibeocr.runtime.recognition.mineru_api.time.sleep", lambda _: None
    )
    server = Server("running")
    with pytest.raises(MineruCancelled):
        MineruApiClient(
            "http://127.0.0.1", transport=httpx.MockTransport(server)
        ).parse(
            [("sample.pdf", b"pdf")],
            MineruConfig(tier=MineruTier.BASIC),
            cancelled=lambda: ("GET", "/v1/parse/jobs/job-1") in server.calls,
        )
    assert not any(
        method == "DELETE" and path.startswith("/v1/files/")
        for method, path in server.calls
    )


def test_upstream_disconnect_does_not_resubmit(monkeypatch):
    monkeypatch.setattr(
        "vibeocr.runtime.recognition.mineru_api.time.sleep", lambda _: None
    )
    server = Server()

    def handler(request):
        if request.url.path == "/v1/parse/jobs/job-1":
            raise httpx.RemoteProtocolError("server restarted", request=request)
        return server(request)

    with pytest.raises(httpx.RemoteProtocolError):
        MineruApiClient(
            "http://127.0.0.1", transport=httpx.MockTransport(handler)
        ).parse([("sample.pdf", b"pdf")], MineruConfig(tier=MineruTier.ADVANCED))
    assert server.calls.count(("POST", "/v1/parse/jobs")) == 1
    assert server.job_payload["tier"] == "advanced"


def test_deadline_before_upload_has_no_side_effects():
    server = Server()
    with pytest.raises(TimeoutError):
        MineruApiClient(
            "http://127.0.0.1", timeout=0, transport=httpx.MockTransport(server)
        ).parse([("sample.pdf", b"pdf")], MineruConfig(tier=MineruTier.BASIC))
    assert server.calls == []


def test_catalog_requires_actual_execution_not_just_installed_packages(monkeypatch):
    from vibeocr.runtime.recognition import mineru_readiness as readiness

    monkeypatch.setattr(readiness.metadata, "version", lambda _: "4.0.2")
    monkeypatch.setattr(readiness, "_observations", {})
    assert all(
        tier["availability"] == "preparation_required"
        for tier in readiness.catalog_payload()["tiers"]
    )
    with pytest.raises(MineruConfigError) as caught:
        readiness.require_ready(MineruTier.BASIC)
    assert caught.value.code == ErrorCode.MINERU_TIER_PREPARATION_REQUIRED
    readiness.mark_executed(MineruTier.BASIC)
    readiness.require_ready(MineruTier.BASIC)
    assert (
        next(t for t in readiness.catalog_payload()["tiers"] if t["id"] == "advanced")[
            "availability"
        ]
        == "preparation_required"
    )
    readiness.mark_failed(MineruTier.BASIC)
    with pytest.raises(MineruConfigError):
        readiness.require_ready(MineruTier.BASIC)


@pytest.mark.parametrize("state", ["completed", "failed", "partial", "canceled"])
def test_recorded_native_job_terminal_states(state, monkeypatch):
    import base64

    monkeypatch.setattr(
        "vibeocr.runtime.recognition.mineru_api.time.sleep", lambda _: None
    )
    recorded = json.loads((FIXTURES / f"job-{state}.json").read_text(encoding="utf-8"))
    job = recorded["job"]
    server = Server()

    def handler(request):
        if request.method == "GET" and request.url.path == "/v1/parse/jobs/job-1":
            return httpx.Response(200, json=job)
        if request.method == "GET" and request.url.path.startswith("/v1/files/"):
            return httpx.Response(
                200,
                content=base64.b64decode(
                    recorded["outputs"][request.url.path.split("/")[3]]
                ),
            )
        return server(request)

    client = MineruApiClient("http://127.0.0.1", transport=httpx.MockTransport(handler))
    inputs = [(file["name"], b"pdf") for file in job["files"]]
    config = MineruConfig(tier=MineruTier(job["tier"]))
    if state == "canceled":
        with pytest.raises(MineruCancelled):
            client.parse(inputs, config)
    else:
        result = client.parse(inputs, config)
        for file in job["files"]:
            expected = (
                MineruDocument if file["status"] == "completed" else MineruApiError
            )
            assert isinstance(result[file["name"]], expected)


def test_confirmed_native_cancel_deletes_owned_inputs(monkeypatch):
    monkeypatch.setattr(
        "vibeocr.runtime.recognition.mineru_api.time.sleep", lambda _: None
    )
    server = Server("running")
    recorded = json.loads((FIXTURES / "cancel-response.json").read_text())

    def handler(request):
        if request.method == "DELETE" and request.url.path == "/v1/parse/jobs/job-1":
            server.calls.append((request.method, request.url.path))
            return httpx.Response(200, json=recorded)
        return server(request)

    with pytest.raises(MineruCancelled):
        MineruApiClient(
            "http://127.0.0.1", transport=httpx.MockTransport(handler)
        ).parse(
            [("sample.pdf", b"pdf")],
            MineruConfig(tier=MineruTier.BASIC),
            cancelled=lambda: ("GET", "/v1/parse/jobs/job-1") in server.calls,
        )
    assert ("DELETE", "/v1/parse/jobs/job-1") in server.calls
    assert ("DELETE", "/v1/files/input-1") in server.calls


def test_native_code_body_is_not_rendered_markdown():
    # Validated docvortex.middle/2.0 and rendered by locked docvortex 0.4.14.
    middle = json.loads((FIXTURES / "code-middle.json").read_text(encoding="utf-8"))
    structured = json.loads(
        (FIXTURES / "code-structured.json").read_text(encoding="utf-8")
    )
    assert structured["pages"][0]["blocks"][0]["content"] == "```python\nprint(1)\n```"
    result = project_document(MineruDocument("", structured, middle, b""))
    assert result.raw_text == "print(1)"
    assert result.content_list[0]["code_body"] == "print(1)"
    assert result.markdown_text == "```\nprint(1)\n```"
    assert "<pre><code>print(1)</code></pre>" in result.html_text


def test_native_semantic_text_lists_equations_and_code_annotations():
    # Validated and rendered with locked docvortex 0.4.14; no OCR inference claim.
    middle = json.loads((FIXTURES / "semantic-middle.json").read_text(encoding="utf-8"))
    structured = json.loads(
        (FIXTURES / "semantic-structured.json").read_text(encoding="utf-8")
    )
    result = project_document(MineruDocument("", structured, middle, b""))
    assert (
        result.raw_text
        == "A * B < C\nfirst\nsecond\nx=1\nExample code\nprint(1)\nCode note"
    )
    assert result.content_list[1]["list_items"] == ["first", "second"]
    assert "- first\n- second" in result.markdown_text
    assert "- - " not in result.markdown_text
    assert "$$x=1$$" in result.markdown_text
    assert "Example code\n\n```\nprint(1)\n```\n\nCode note" in result.markdown_text
    assert "A \\* B &lt; C" in result.markdown_text
    assert "<p>A * B &lt; C</p>" in result.html_text
    assert "<li>first</li><li>second</li>" in result.html_text
    assert "Example code" in result.html_text and "Code note" in result.html_text


@pytest.mark.parametrize("suffix", ["docx", "pdf"])
def test_native_visual_annotations_survive_all_projections(suffix):
    # Validated/rendered by docvortex 0.4.14; synthetic schema fixture, not OCR.
    middle = json.loads(
        (FIXTURES / f"annotations-{suffix}-middle.json").read_text(encoding="utf-8")
    )
    structured = json.loads(
        (FIXTURES / f"annotations-{suffix}-structured.json").read_text(encoding="utf-8")
    )
    result = project_document(MineruDocument("", structured, middle, b""))
    for kind in ("image", "chart", "table"):
        caption = f"{kind} caption *literal*"
        footnote = f"{kind} note <b>literal</b>"
        assert result.raw_text.count(caption) == 1
        assert result.raw_text.count(footnote) == 1
        assert result.raw_text.index(caption) < result.raw_text.index(footnote)
        assert f"{kind} caption \\*literal\\*" in result.markdown_text
        assert f"{kind} note &lt;b&gt;literal&lt;/b&gt;" in result.markdown_text
        assert caption in result.html_text
        assert f"{kind} note &lt;b&gt;literal&lt;/b&gt;" in result.html_text
    assert result.raw_text.count("Cell") == 1


@pytest.mark.parametrize("suffix", ["docx", "pdf"])
def test_native_visual_bodies_formula_fallback_and_nested_code_fence(suffix):
    middle = json.loads(
        (FIXTURES / f"annotations-{suffix}-middle.json").read_text(encoding="utf-8")
    )
    structured = json.loads(
        (FIXTURES / f"annotations-{suffix}-structured.json").read_text(encoding="utf-8")
    )
    result = project_document(MineruDocument("", structured, middle, b""))
    assert result.raw_text.count("image body *literal*") == 1
    assert "image body \\*literal\\*" in result.markdown_text
    assert "image body *literal*" in result.html_text
    for projection in (result.raw_text, result.markdown_text, result.html_text):
        assert "Revenue" in projection and "42" in projection
    assert "<table>" not in result.raw_text
    equation = next(b for b in result.content_list if b["type"] == "equation")
    assert equation["img_path"] in result.images
    assert f"![]({equation['img_path']})" in result.markdown_text
    assert f'<img src="{equation["img_path"]}"' in result.html_text
    assert "````\nfirst\n```\n# literal heading\nlast\n````" in result.markdown_text
    assert (
        "<pre><code>first\n```\n# literal heading\nlast</code></pre>"
        in result.html_text
    )


@pytest.mark.parametrize("suffix", ["docx", "pdf"])
def test_native_table_fallback_and_index_keep_semantic_content(suffix):
    middle = json.loads(
        (FIXTURES / f"annotations-{suffix}-middle.json").read_text(encoding="utf-8")
    )
    structured = json.loads(
        (FIXTURES / f"annotations-{suffix}-structured.json").read_text(encoding="utf-8")
    )
    result = project_document(MineruDocument("", structured, middle, b""))
    fallback = next(b for b in result.content_list if b["type"] == "table_unparsed")
    assert fallback["text"] == "Item  Count\nApple  12"
    for projection in (result.raw_text, result.markdown_text, result.html_text):
        assert projection.count("Fallback caption") == 1
        assert projection.count("Fallback note") == 1
        assert "Apple  12" in projection
    assert "Fallback caption\nItem  Count\nApple  12\nFallback note" in result.raw_text
    assert "Chapter A ... 1\nChapter B ... 2" in result.raw_text
    assert "Chapter A ... 1<br>Chapter B ... 2" in result.html_text

    image_table = [b for b in result.content_list if b["type"] == "table_unparsed"][1]
    assert image_table["text"] == ""
    assert image_table["img_path"] in result.images
    assert image_table["img_path"] in result.markdown_text
    assert image_table["img_path"] in result.html_text
    for projection in (result.raw_text, result.markdown_text, result.html_text):
        assert projection.count("Image fallback caption") == 1
        assert projection.count("Image fallback note") == 1


@pytest.mark.parametrize("suffix", ["docx", "pdf"])
def test_native_empty_code_keeps_annotations(suffix):
    middle = json.loads(
        (FIXTURES / f"annotations-{suffix}-middle.json").read_text(encoding="utf-8")
    )
    structured = json.loads(
        (FIXTURES / f"annotations-{suffix}-structured.json").read_text(encoding="utf-8")
    )
    result = project_document(MineruDocument("", structured, middle, b""))
    for projection in (result.raw_text, result.markdown_text, result.html_text):
        assert projection.count("Empty code caption") == 1
        assert projection.count("Empty code note") == 1
    assert "Empty code caption\n\nEmpty code note" in result.markdown_text
    assert "<p>Empty code caption</p>\n<p>Empty code note</p>" in result.html_text


@pytest.mark.parametrize("suffix", ["docx", "pdf"])
@pytest.mark.parametrize("fmt", ["docx", "xlsx"])
def test_native_projection_survives_office_export(suffix, fmt, tmp_path):
    from vibeocr.runtime.documents.export_service import ExportService

    middle = json.loads(
        (FIXTURES / f"annotations-{suffix}-middle.json").read_text(encoding="utf-8")
    )
    structured = json.loads(
        (FIXTURES / f"annotations-{suffix}-structured.json").read_text(encoding="utf-8")
    )
    result = project_document(MineruDocument("", structured, middle, b""))
    path = tmp_path / f"result.{fmt}"
    assert ExportService.export(result, path, fmt)
    if fmt == "docx":
        from docx import Document

        document = Document(path)
        text = "\n".join(p.text for p in document.paragraphs)
        assert len(document.inline_shapes) == 2
    else:
        from openpyxl import load_workbook

        workbook = load_workbook(path)
        text = "\n".join(
            str(value)
            for sheet in workbook
            for row in sheet.values
            for value in row
            if value is not None
        )
        workbook.close()
    for value in (
        "Fallback caption",
        "Apple  12",
        "Fallback note",
        "Image fallback caption",
        "Image fallback note",
        "Empty code caption",
        "Empty code note",
        "image caption *literal*",
        "image body *literal*",
        "image note <b>literal</b>",
        "chart caption *literal*",
        "Revenue",
        "42",
        "chart note <b>literal</b>",
        "Chapter A ... 1\nChapter B ... 2",
        "Reference Alpha 2026",
    ):
        assert text.count(value) == 1
