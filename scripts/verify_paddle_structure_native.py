"""Native retained Paddle graph evidence on isolated, pre-provisioned ORT environments."""

from __future__ import annotations

import argparse
import copy
import dataclasses
import importlib.metadata
import json
import math
import os
import socket
import statistics
import subprocess
import time
from pathlib import Path


def _json(value):
    if hasattr(value, "tolist"):
        return value.tolist()
    if isinstance(value, bytes):
        return {"bytes": len(value)}
    if hasattr(value, "size") and hasattr(value, "mode"):
        return {"image_size": value.size, "mode": value.mode}
    raise TypeError(type(value).__name__)


def _check_exports(result, destination):
    from docx import Document
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from vibeocr.runtime.documents.export_service import ExportService
    from vibeocr.runtime.documents.tables.blocks import table_model_from_block
    from vibeocr.runtime.documents.tables.html_adapter import table_model_from_html

    tables = [
        table_model_from_block(b)
        for b in result.content_list
        if b.get("type") == "table"
    ]
    for fmt in ("html", "xlsx", "docx"):
        assert ExportService.export(result, destination.with_suffix("." + fmt), fmt)
    workbook = load_workbook(destination.with_suffix(".xlsx"))
    doc = Document(destination.with_suffix(".docx"))
    assert len(doc.tables) == len(tables)
    for index, table in enumerate(tables):
        sheet = workbook[f"表格 {index + 1}"]
        merges = set()
        for cell in table.cells:
            actual = sheet.cell(cell.row + 1, cell.column + 1).value or ""
            assert actual == cell.text, ("xlsx text", cell, actual)
            assert doc.tables[index].cell(cell.row, cell.column).text == cell.text
            if cell.rowspan > 1 or cell.colspan > 1:
                merges.add(
                    f"{get_column_letter(cell.column + 1)}{cell.row + 1}:{get_column_letter(cell.column + cell.colspan)}{cell.row + cell.rowspan}"
                )
                assert (
                    doc.tables[index].cell(cell.row, cell.column)._tc
                    is doc.tables[index]
                    .cell(cell.row + cell.rowspan - 1, cell.column + cell.colspan - 1)
                    ._tc
                )
        assert {str(value) for value in sheet.merged_cells.ranges} == merges
    if len(tables) == 1:
        actual = table_model_from_html(
            destination.with_suffix(".html").read_text(encoding="utf-8"),
            table_id=tables[0].table_id,
        )

        def cells(table):
            return [
                (c.row, c.column, c.rowspan, c.colspan, c.text) for c in table.cells
            ]

        assert cells(actual) == cells(tables[0])
    workbook.close()


def _check_result(name, result, width, height):
    tables = [b["table"] for b in result.content_list if b.get("type") == "table"]
    assert result.content_list, "empty structured result"
    assert all(b.get("type") != "seal" for b in result.content_list)
    for block in result.content_list:
        box = block.get("bbox")
        if box is not None:
            assert len(box) == 4 and all(math.isfinite(v) for v in box)
            assert 0 <= box[0] <= box[2] <= width and 0 <= box[1] <= box[3] <= height, (
                box
            )
    if name in ("table_merged_zh_en_90.png", "table_merged_zh_en_270.png"):
        assert result.preproc_angle in (90, 270), (
            "expected quarter-turn correction",
            result.preproc_angle,
        )
        assert result.preproc_img_w > result.preproc_img_h > 0, (
            "expected corrected landscape space",
            result.preproc_img_w,
            result.preproc_img_h,
        )
        from vibeocr.runtime.recognition.ocr_service import OCRService

        normalized = copy.deepcopy(result)
        OCRService._normalize_result_bbox(normalized, image=None)
        assert (normalized.image_width, normalized.image_height) == (width, height)
        for raw, mapped in zip(
            result.content_list, normalized.content_list, strict=True
        ):
            if raw.get("bbox") is not None:
                expected = [
                    value / (width if index % 2 == 0 else height) * 1000
                    for index, value in enumerate(raw["bbox"])
                ]
                assert mapped["bbox"] == expected, (
                    "normalized source space",
                    mapped,
                    expected,
                )
                assert all(0 <= value <= 1000 for value in mapped["bbox"])
            assert mapped.get("table") == raw.get("table"), (
                "canonical cell coordinates changed"
            )
    if "no_table" in name:
        assert not tables, "ordinary prose became a table"
    elif "table_merged" in name:
        assert len(tables) == 1, f"expected one table, got {len(tables)}"
        table = tables[0]
        assert (table["row_count"], table["column_count"]) == (5, 4), (
            "expected 5x4",
            table["row_count"],
            table["column_count"],
        )
        cells = {(c["row"], c["column"]): c for c in table["cells"]}
        assert cells[0, 0]["colspan"] == 4
        assert cells[2, 0]["rowspan"] == 2 and "Merged" in cells[2, 0]["text"]
        for location, text in {
            (2, 1): "3",
            (2, 2): "12.50",
            (3, 1): "7",
            (3, 2): "0.99",
            (4, 1): "10",
            (4, 2): "45.47",
        }.items():
            assert cells[location]["text"] == text, (location, cells[location])
    elif name == "wireless_dispatch.png":
        assert len(tables) == 1, ("one dispatch table expected", len(tables))
        table = tables[0]
        assert (table["row_count"], table["column_count"]) == (5, 5), (
            "expected 5x5",
            table["row_count"],
            table["column_count"],
        )
        expected = [
            ["Code", "Item", "Qty", "Unit", "Total"],
            ["A17", "Bolts", "6", "2.25", "13.50"],
            ["B23", "Washers", "12", "0.40", "4.80"],
            ["C08", "Hinges", "3", "8.50", "25.50"],
            ["D42", "Cable", "8", "1.75", "14.00"],
        ]
        cells = {(cell["row"], cell["column"]): cell for cell in table["cells"]}
        for row, values in enumerate(expected):
            for column, value in enumerate(values):
                cell = cells[row, column]
                assert cell["rowspan"] == cell["colspan"] == 1, (row, column, cell)
                assert " ".join(cell["text"].casefold().split()) == value.casefold(), (
                    row,
                    column,
                    cell["text"],
                    value,
                )
        assert "WAREHOUSE DISPATCH" in result.raw_text, "dispatch title missing"
    elif "table_wireless" in name:
        assert len(tables) == 1, f"expected one table, got {len(tables)}"
        table = tables[0]
        cells = {(c["row"], c["column"]): c["text"] for c in table["cells"]}
        if "geometry_holdout" in name:
            assert (table["row_count"], table["column_count"]) == (6, 3), (
                "expected 6x3",
                table["row_count"],
                table["column_count"],
            )
            expected = [
                ["Material", "Qty", "Shelf"],
                ["Copper wire", "12", "A3"],
                ["Steel bolts", "48", "B2"],
                ["Rubber pads", "16", "C5"],
                ["Glass tubes", "9", "D1"],
                ["Paper rolls", "25", "E4"],
            ]
            for row, values in enumerate(expected):
                for column, value in enumerate(values):
                    assert cells[row, column] == value, (
                        row,
                        column,
                        cells[row, column],
                        value,
                    )
        elif "acceptance" in name:
            assert (table["row_count"], table["column_count"]) == (5, 4), (
                "expected 5x4",
                table["row_count"],
                table["column_count"],
            )
            expected = [
                ["Code", "Count", "Price", "Total"],
                ["A100", "2", "8.50", "17.00"],
                ["B200", "6", "3.00", "18.00"],
                ["C300", "4", "5.25", "21.00"],
                ["D400", "1", "9.00", "9.00"],
            ]
            for row, values in enumerate(expected):
                for column, value in enumerate(values):
                    assert cells[row, column] == value, (
                        row,
                        column,
                        cells[row, column],
                    )
        elif "final" in name:
            assert (table["row_count"], table["column_count"]) == (5, 2), (
                "expected 5x2",
                table["row_count"],
                table["column_count"],
            )
            expected = [
                ["Item", "Count"],
                ["Screws", "24"],
                ["Clips", "18"],
                ["Pins", "36"],
                ["Boxes", "5"],
            ]
            for row, values in enumerate(expected):
                for column, value in enumerate(values):
                    assert cells[row, column] == value, (
                        row,
                        column,
                        cells[row, column],
                    )
        elif "holdout" in name:
            assert (table["row_count"], table["column_count"]) == (4, 3), (
                "expected 4x3",
                table["row_count"],
                table["column_count"],
            )
            for location, text in {
                (1, 0): "Notebook",
                (1, 1): "4",
                (1, 2): "28.00",
                (2, 0): "Pencil",
                (2, 2): "6.00",
                (3, 0): "Folder",
                (3, 2): "15.00",
            }.items():
                assert cells[location] == text, (location, cells[location])
        else:
            assert table["column_count"] == 4, (
                f"expected four wireless columns, got {table['column_count']}"
            )
            assert any("Merged" in value for value in cells.values())
            assert any(
                "GOAL103" in str(b.get("text", "")) for b in result.content_list
            ) or "GOAL103" in str(table)
            row = next(
                row
                for (row, column), text in cells.items()
                if column == 1 and text == "3"
            )
            assert (
                cells[row, 2] == "12.50"
                and cells[row + 1, 1] == "7"
                and cells[row + 1, 2] == "0.99"
            )
    elif "document_mixed" in name:
        assert tables
        assert {"text", "image", "formula", "footer"} <= {
            b["type"] for b in result.content_list
        }, f"mixed content types: {[b['type'] for b in result.content_list]}"
    elif "document_columns" in name:
        text = "\n".join(str(b.get("text", "")) for b in result.content_list)
        markers = [
            "DOCUMENT COLUMN ORDER",
            "LEFT FIRST",
            "LEFT SECOND",
            "RIGHT FIRST",
            "RIGHT SECOND",
            "SPANNING FOOTER END OF DOCUMENT",
        ]
        offsets = [text.index(marker) for marker in markers]
        assert offsets == sorted(offsets)
        assert "SPANNING FOOTER END OF DOCUMENT" in result.raw_text


def _construction_only_role(role, *, wireless_e2e, wireless_cell_html):
    # Only the two explicitly selected, mutually exclusive public routes qualify.
    # Other combinations fail closed instead of guessing which model executed.
    if wireless_cell_html and not wireless_e2e:
        return role.endswith(".wireless_table_rec_model")
    if wireless_e2e and not wireless_cell_html:
        return role.endswith(".wireless_table_cells_detection_model")
    return False


def _runner_roles(pipeline):
    roles = {}
    seen = set()

    def visit(node, path):
        while hasattr(node, "_pipeline"):
            node = node._pipeline
        if id(node) in seen:
            return
        seen.add(id(node))
        runner = getattr(node, "runner", None)
        if runner is not None:
            roles[id(runner)] = path
            return
        for name, child in vars(node).items():
            if hasattr(child, "__dict__") and (
                name.endswith("_model") or name.endswith("_pipeline")
            ):
                visit(child, path + "." + name)

    visit(pipeline.paddlex_pipeline, "root")
    return roles


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--fixtures", type=Path, required=True)
    parser.add_argument(
        "--source", choices=("modelscope", "huggingface"), required=True
    )
    parser.add_argument("--device", choices=("cpu", "cuda", "directml"), required=True)
    parser.add_argument(
        "--preset", choices=("tiny", "small", "medium"), default="small"
    )
    parser.add_argument(
        "--mode", choices=("TABLE_RECOGNITION", "PP-StructureV3"), required=True
    )
    parser.add_argument("--local-only", action="store_true")
    parser.add_argument("--all-aux", action="store_true")
    parser.add_argument("--warm", type=int, default=5)
    parser.add_argument(
        "--wireless-e2e",
        action="store_true",
        help="Diagnose the internal TABLE_RECOGNITION E2E route; this is not a public wire/UI option",
    )
    parser.add_argument("--input", action="append")
    args = parser.parse_args()
    if args.wireless_e2e and args.mode != "TABLE_RECOGNITION":
        parser.error("--wireless-e2e is only an internal independent-table diagnostic")
    if args.output.exists():
        raise FileExistsError(args.output)
    args.output.mkdir(parents=True)
    if not os.environ.get("VIBEOCR_SHARED_MODEL_CACHE"):
        raise RuntimeError("Set VIBEOCR_SHARED_MODEL_CACHE to the isolated cache")
    os.environ.update(
        VIBEOCR_MANAGED_ENVIRONMENT_RECIPE=f"paddleocr-ort-{args.device}",
        VIBEOCR_MANAGED_ENVIRONMENT_ID="retained-smoke",
        VIBEOCR_MANAGED_ENVIRONMENT_REVISION="1",
        PADDLE_PDX_MODEL_SOURCE=args.source,
        PADDLE_PDX_CACHE_HOME=str(args.output / "private"),
        PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK="True",
    )
    blocked_connections = []
    original_connect = socket.socket.connect
    if args.local_only:

        def deny_connect(_socket, address):
            blocked_connections.append(str(address))
            raise OSError("Network disabled for offline retained-graph verification")

        socket.socket.connect = deny_connect
    started = time.perf_counter()
    import onnxruntime
    import psutil
    from PIL import Image
    from vibeocr.runtime.recognition import paddle_onnx
    from vibeocr.runtime.recognition.core.pipelines import (
        pipeline_pp_structure,
        pipeline_table,
    )

    installed = {
        dist.metadata["Name"].lower() for dist in importlib.metadata.distributions()
    }
    assert not installed & {"paddlepaddle", "paddlepaddle-gpu", "torch", "transformers"}
    native_session = onnxruntime.InferenceSession

    def profiled_session(*a, **kw):
        kw["sess_options"].enable_profiling = True
        kw["sess_options"].profile_file_prefix = str(
            args.output / f"ort-{time.time_ns()}"
        )
        return native_session(*a, **kw)

    onnxruntime.InferenceSession = profiled_session
    if args.mode == "TABLE_RECOGNITION":
        options = pipeline_table.TableRecognitionOptions(
            model_preset=args.preset,
            use_doc_unwarping=args.all_aux,
            use_e2e_wireless_table_rec_model=args.wireless_e2e,
            use_wireless_table_cells_trans_to_html=False,
        )
        recognize = pipeline_table._recognize_table
        names = [
            "table_merged_zh_en.png",
            "table_wireless.png",
            "table_wireless_holdout.png",
            "text_no_table.png",
            "document_no_table.png",
        ]
    else:
        options = pipeline_pp_structure.PPStructureV3Options(
            model_preset=args.preset,
            use_doc_unwarping=args.all_aux,
            use_textline_orientation=args.all_aux,
        )
        recognize = pipeline_pp_structure._recognize_pp_structure
        names = [
            "document_mixed.png",
            "document_columns.png",
            "document_columns_90.png",
            "table_wireless_holdout.png",
        ]
    snapshot = paddle_onnx.freeze_request(vars(options), pipeline_id=args.mode)
    report = {
        "source_sha": subprocess.check_output(
            ["git", "rev-parse", "HEAD"], text=True
        ).strip(),
        "tracked_worktree_dirty": subprocess.run(
            ["git", "diff", "--quiet", "HEAD"], check=False
        ).returncode
        != 0,
        "snapshot": snapshot,
        "versions": {
            name: importlib.metadata.version(name)
            for name in (
                "paddlex",
                "paddleocr",
                "onnxruntime-directml" if args.device == "directml" else "onnxruntime",
            )
        },
        "request_options": vars(options),
        "network_disabled": args.local_only,
        "inputs": [],
        "warm_seconds": [],
        "profiles": [],
        "passed": False,
    }
    pipeline = None
    try:
        pipeline = paddle_onnx.create_pipeline(
            json.dumps(snapshot), local_models_only=args.local_only
        )
        report["import_construct_seconds"] = time.perf_counter() - started
        report["sessions"] = pipeline.sessions

        class Service:
            def get_or_create_pipeline(self, *_a, **_kw):
                return pipeline

        for index, name in enumerate(args.input or names):
            image_path = args.fixtures / name
            begin = time.perf_counter()
            result = recognize(Service(), str(image_path), options)
            elapsed = time.perf_counter() - begin
            if index == 0:
                report["cold_seconds"] = time.perf_counter() - started
            result_file = args.output / (name + ".json")
            result_file.write_text(
                json.dumps(
                    dataclasses.asdict(result),
                    default=_json,
                    ensure_ascii=False,
                    indent=2,
                ),
                encoding="utf-8",
            )
            with Image.open(image_path) as image:
                width, height = image.size
            if result.preproc_img_w and result.preproc_img_h:
                width, height = result.preproc_img_w, result.preproc_img_h
            input_report = {
                "name": name,
                "seconds": elapsed,
                "blocks": len(result.content_list),
                "passed": False,
            }
            try:
                _check_result(name, result, width, height)
                _check_exports(result, args.output / name)
                input_report["passed"] = True
            except (AssertionError, ValueError) as exc:
                input_report["failure"] = f"{type(exc).__name__}: {exc}"
            report["inputs"].append(input_report)
            if index == 0:
                for _ in range(args.warm):
                    begin = time.perf_counter()
                    recognize(Service(), str(image_path), options)
                    report["warm_seconds"].append(time.perf_counter() - begin)
        if report["warm_seconds"]:
            report["warm_median_seconds"] = statistics.median(report["warm_seconds"])
        report["quality_passed"] = all(item["passed"] for item in report["inputs"])
    except Exception as exc:
        report["failure"] = f"{type(exc).__name__}: {exc}"
        raise
    finally:
        profile_errors = []
        if pipeline is not None:
            try:
                expected_provider = {
                    "cpu": "CPUExecutionProvider",
                    "cuda": "CUDAExecutionProvider",
                    "directml": "DmlExecutionProvider",
                }[args.device]
                roles = _runner_roles(pipeline.pipeline)
                for role_index, runner in enumerate(
                    paddle_onnx.owned_runners(pipeline.pipeline)
                ):
                    try:
                        role = roles[id(runner)]
                        construction_dependency = _construction_only_role(
                            role,
                            wireless_e2e=getattr(
                                options, "use_e2e_wireless_table_rec_model", False
                            ),
                            wireless_cell_html=getattr(
                                options, "use_wireless_table_cells_trans_to_html", False
                            ),
                        )
                        profile = runner.session.end_profiling()
                        events = json.loads(Path(profile).read_text(encoding="utf-8"))
                        partitions = {}
                        for event in events:
                            provider = event.get("args", {}).get("provider")
                            if provider:
                                key = provider + ":" + event["args"].get("op_name", "")
                                partitions[key] = partitions.get(key, 0) + 1
                        executed = any(
                            key.startswith(expected_provider + ":")
                            for key in partitions
                        )
                        report["profiles"].append(
                            {
                                "role_index": role_index,
                                "role": role,
                                "execution_requirement": "upstream_constructed_selected_route_not_consumed"
                                if construction_dependency
                                else "required",
                                "model": str(runner.model_dir),
                                "path": profile,
                                "operator_calls": partitions,
                                "target_provider_executed": executed,
                            }
                        )
                        if not executed and not construction_dependency:
                            profile_errors.append(
                                f"role {role_index} {runner.model_dir.name}: no {expected_provider} execution"
                            )
                    except Exception as exc:
                        profile_errors.append(
                            f"role {role_index}: {type(exc).__name__}: {exc}"
                        )
            finally:
                pipeline.close()
        report["profile_errors"] = profile_errors
        report["passed"] = bool(report.get("quality_passed")) and not profile_errors
        report["peak_working_set_bytes"] = getattr(
            psutil.Process().memory_info(), "peak_wset", None
        )
        report["blocked_connections"] = blocked_connections
        (args.output / "report.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        onnxruntime.InferenceSession = native_session
        socket.socket.connect = original_connect
        report["blocked_connections"] = blocked_connections
    if not report["passed"]:
        raise RuntimeError("Retained graph verification failed; see report.json")
    print(
        json.dumps(
            {"passed": report["passed"], "output": str(args.output)}, ensure_ascii=False
        )
    )


if __name__ == "__main__":
    main()
